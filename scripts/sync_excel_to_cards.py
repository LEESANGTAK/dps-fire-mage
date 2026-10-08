#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
sync_excel_to_cards.py
- src/data/roguelike-cards.xlsx 파일을 읽어 게임 엔진용 src/data/roguelike-cards.json 파일로 변환/동기화합니다.
- 기획자가 엑셀에서 수치, 텍스트, 드롭다운 옵션을 수정한 후 본 스크립트를 실행하면 인게임에 즉시 반영됩니다.
- 열은 위치가 아니라 1행 헤더 이름(scripts/card_columns.py)으로 찾습니다.
  열 구성이 다른 구형 엑셀이면 JSON을 덮어쓰지 않고 중단합니다.
"""

import json
import os
import sys

if sys.stdout.encoding != 'utf-8':
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except Exception:
        pass

import openpyxl

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from card_columns import CARD_COLUMNS, LEGACY_HEADERS  # noqa: E402

ROOT_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
JSON_PATH = os.path.join(ROOT_DIR, 'src', 'data', 'roguelike-cards.json')
EXCEL_PATH = os.path.join(ROOT_DIR, 'src', 'data', 'roguelike-cards.xlsx')
DECK_FALLBACK_SCRIPT = os.path.join(ROOT_DIR, 'scripts', 'sync-card-fallback.cjs')


def parse_number(val):
    num = float(val)
    return int(num) if num.is_integer() else num


def sync_excel_to_json():
    if not os.path.exists(EXCEL_PATH):
        print(f"❌ [에러] 엑셀 파일을 찾을 수 없습니다: {EXCEL_PATH}")
        sys.exit(1)

    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
    ws = wb['RoguelikeCards'] if 'RoguelikeCards' in wb.sheetnames else wb.active

    header_cells = [str(c.value).strip() if c.value is not None else '' for c in ws[1]]
    legacy = [h for h in header_cells if h in LEGACY_HEADERS]
    missing = [name for name, _, _, _ in CARD_COLUMNS if name not in header_cells]
    if legacy or missing:
        print("❌ [중단] 엑셀 열 구성이 현재 카드 형식과 다릅니다. JSON은 변경하지 않았습니다.")
        if legacy:
            print(f"   구형 열: {', '.join(legacy)}")
        if missing:
            print(f"   없는 열: {', '.join(missing)}")
        print("   먼저 `npm run export-cards`로 엑셀을 현재 형식으로 다시 생성한 뒤 수정해 주세요.")
        sys.exit(1)
    col_of = {key: header_cells.index(name) + 1 for name, key, _, _ in CARD_COLUMNS}

    def cell(row_idx, key):
        return ws.cell(row=row_idx, column=col_of[key]).value

    cards = []
    for row_idx in range(2, ws.max_row + 1):
        card_id = cell(row_idx, 'id')
        if not card_id or not str(card_id).strip():
            continue

        try:
            max_stack = int(cell(row_idx, 'maxStack')) if cell(row_idx, 'maxStack') is not None else 1
        except (ValueError, TypeError):
            max_stack = 1

        card = {
            "id": str(card_id).strip(),
            "title": str(cell(row_idx, 'title') or '').strip(),
            "icon": str(cell(row_idx, 'icon') or '').strip(),
            "rarity": str(cell(row_idx, 'rarity') or 'common').strip(),
            "build": str(cell(row_idx, 'build') or 'rapid').strip(),
            "maxStack": max_stack,
            "desc": str(cell(row_idx, 'desc') or '').strip()
        }

        stats = {}
        for _, key, _, _ in CARD_COLUMNS:
            if not key.startswith('stats.'):
                continue
            val = cell(row_idx, key)
            if val is not None and str(val).strip() != '':
                try:
                    stats[key[len('stats.'):]] = parse_number(val)
                except (ValueError, TypeError):
                    pass
        if stats:
            card["stats"] = stats

        for key in ('prerequisite', 'spellKey', 'mechanic'):
            val = cell(row_idx, key)
            if val and str(val).strip():
                card[key] = str(val).strip()

        params = cell(row_idx, 'mechanicParams')
        if params and str(params).strip():
            try:
                card["mechanicParams"] = json.loads(str(params).strip())
            except json.JSONDecodeError:
                print(f"⚠️ [경고] {card['id']}의 기믹파라미터 JSON 형식이 올바르지 않아 건너뜁니다: {params}")

        cards.append(card)

    with open(JSON_PATH, 'w', encoding='utf-8') as f:
        json.dump(cards, f, ensure_ascii=False, indent=2)
        f.write('\n')

    # file:// 실행용 내장 fallback 카드도 함께 갱신
    if os.system(f'node "{DECK_FALLBACK_SCRIPT}"') != 0:
        print("⚠️ [경고] 내장 fallback 카드 동기화에 실패했습니다. `node scripts/sync-card-fallback.cjs`를 직접 실행해 주세요.")

    print(f"✨ [동기화 성공] 엑셀({EXCEL_PATH}) ➔ JSON({JSON_PATH})")
    print(f"   총 {len(cards)}종의 파워업 카드가 업데이트되었습니다.")
    print("   브라우저에서 새로고침(F5)을 누르면 즉시 게임에 반영됩니다.")


if __name__ == "__main__":
    sync_excel_to_json()
